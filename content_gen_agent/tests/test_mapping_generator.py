"""CG-016: the tables joining questions, errors and misconceptions.

Two of the three are derived rather than generated, and the tests that matter
most are the ones checking that derivation against the platform's own data.
If our rule for building misconception_errors disagrees with what the platform
holds, that is worth knowing before we ship 46 misconceptions built on it.

No network.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from llm_client import FakeLLMClient, LLMError      # noqa: E402
from mapping_generator import (                     # noqa: E402
    DEFAULT_CONFIDENCE,
    ERROR_MAP_SYSTEM_PROMPT,
    MappingError,
    build_error_map_prompt,
    build_misconception_errors,
    option_texts,
    unmapped_by_question,
    direct_failures,
    generate_question_error_map,
)
from models import (                                # noqa: E402
    AnswerSpecRow,
    AnswerType,
    DetectionMethod,
    ErrorTypeRow,
    MisconceptionRow,
    QuestionRow,
    QuestionStatus,
    QuestionType,
    RelationshipType,
    Severity as ErrSeverity,
    VerificationMethod,
)
from sources import SCHEMA_TEMPLATE                 # noqa: E402

needs_template = pytest.mark.skipif(
    SCHEMA_TEMPLATE is None, reason="schema template not available",
)


def _error(code, skill="T01.M2", name="Error"):
    return ErrorTypeRow(
        error_code=code, error_name=name,
        description="Student writes the wrong thing.",
        related_micro_skill_id=skill, severity=ErrSeverity.HIGH,
        detection_method=DetectionMethod.SYMBOLIC_PATTERN, active=True,
    )


def _misconception(slug, rule):
    return MisconceptionRow(
        misconception_id=f"MIS-T01-{slug}", name=f"Belief {slug}",
        description="Student believes the wrong thing.",
        diagnosis_rule=rule, active=True, version="1.0",
    )


def _question(i=1, qtype=QuestionType.SINGLE_CHOICE, text=None):
    return QuestionRow(
        question_id=f"Q-T01-{i:03d}", topic_id="ALG-ORI-01",
        question_text=text or "Which rule works? a) n+5 b) 5n c) n-5",
        question_type=qtype, difficulty=2,
        answer_spec_id=f"ANS-T01-{i:03d}", item_family_id="FAM-T01-X",
        source_provenance_id="SRC-NABLIX-T01-001",
        status=QuestionStatus.GENERATED, version="1.0",
    )


def _answer(i=1, wrong="5n | n-5"):
    return AnswerSpecRow(
        answer_spec_id=f"ANS-T01-{i:03d}", question_id=f"Q-T01-{i:03d}",
        answer_type=AnswerType.ALGEBRAIC_EXPRESSION, canonical_answer="n + 5",
        accepted_answers="n+5", common_wrong_answers=wrong,
        verification_method=VerificationMethod.SYMBOLIC_EQUIVALENCE,
        required_units=None, explanation_required=False,
        answer_steps="1. Do it.",
    )


# ──────────────────────────────────────────────────────────────────────
# Derived: misconception -> error
# ──────────────────────────────────────────────────────────────────────

def test_the_rule_that_names_an_error_produces_the_mapping():
    errors = [_error("ERR-T01-ADD-AS-MULTIPLY")]
    rows, issues = build_misconception_errors(
        [_misconception("ADD", "Trigger after ERR-T01-ADD-AS-MULTIPLY when repeated.")],
        errors,
    )
    assert len(rows) == 1
    assert rows[0].error_code == "ERR-T01-ADD-AS-MULTIPLY"
    assert rows[0].confidence_weight == DEFAULT_CONFIDENCE == 1.0
    assert not [i for i in issues if i.is_error]


def test_a_rule_naming_two_errors_produces_two_rows():
    errors = [_error("ERR-T01-A"), _error("ERR-T01-B")]
    rows, _ = build_misconception_errors(
        [_misconception("BOTH", "Trigger after ERR-T01-A or ERR-T01-B when repeated.")],
        errors,
    )
    assert {r.error_code for r in rows} == {"ERR-T01-A", "ERR-T01-B"}


def test_the_same_code_named_twice_produces_one_row():
    errors = [_error("ERR-T01-A")]
    rows, _ = build_misconception_errors(
        [_misconception("X", "Trigger after ERR-T01-A, and again after ERR-T01-A.")],
        errors,
    )
    assert len(rows) == 1


def test_an_error_no_misconception_explains_warns():
    """The review: every error must resolve to a misconception, or the tutor
    can detect something it has nothing to re-teach for."""
    errors = [_error("ERR-T01-A"), _error("ERR-T01-ORPHAN")]
    _, issues = build_misconception_errors(
        [_misconception("X", "Trigger after ERR-T01-A when repeated.")], errors,
    )
    warning = next(i for i in issues if "resolve to no misconception" in i.message)
    assert not warning.is_error
    assert "ERR-T01-ORPHAN" in warning.message


def test_a_rule_naming_nothing_real_is_an_error():
    """CG-015 refuses this, so reaching it means the two ran against
    different error sets, which is worth shouting about."""
    _, issues = build_misconception_errors(
        [_misconception("X", "Trigger after ERR-T01-GHOST when repeated.")],
        [_error("ERR-T01-A")],
    )
    assert any(i.is_error and "nothing can trigger it" in i.message for i in issues)


@needs_template
def test_our_derivation_agrees_with_the_platform_data():
    """The test that decides whether deriving is safe at all.

    Rebuilds misconception_errors from the template's own diagnosis_rules and
    compares with the table the platform actually holds.
    """
    from collections import Counter

    from openpyxl import load_workbook
    wb = load_workbook(SCHEMA_TEMPLATE)

    def rows(sheet):
        ws = wb[sheet]
        header = [c.value for c in ws[1] if c.value is not None]
        return [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)
                if any(v is not None for v in r)]

    errors = [
        ErrorTypeRow(
            error_code=r["error_code"], error_name=r["error_name"],
            description=r["description"],
            related_micro_skill_id=r["related_micro_skill_id"],
            severity=ErrSeverity(r["severity"]),
            detection_method=DetectionMethod(r["detection_method"]),
            active=bool(r["active"]),
        )
        for r in rows("error_types")
    ]
    misconceptions = [
        MisconceptionRow(
            misconception_id=r["misconception_id"], name=r["name"],
            description=r["description"], diagnosis_rule=r["diagnosis_rule"],
            active=bool(r["active"]), version=str(r["version"]),
        )
        for r in rows("misconceptions")
    ]

    derived, _ = build_misconception_errors(misconceptions, errors)
    ours = {(r.misconception_id, r.error_code) for r in derived}
    theirs = {(r["misconception_id"], r["error_code"])
              for r in rows("misconception_errors")}

    # Nothing we derive should be absent from the platform's table. The
    # reverse is allowed: it holds two extra pairs, both cross-topic
    # (MIS-T03 to ERR-T01), which is the same legacy id mess recorded in M2.
    assert ours - theirs == set(), f"we derive pairs the platform does not have: {ours - theirs}"
    assert len(ours) >= len(theirs) - 2


# ──────────────────────────────────────────────────────────────────────
# Derived: misconception -> skills
# ──────────────────────────────────────────────────────────────────────

def test_direct_failures_come_from_the_errors_own_skills():
    """Not a judgement: a misconception causes errors, and each error already
    records the skill it shows a gap in."""
    errors = [_error("ERR-T01-A", skill="T01.M2"),
              _error("ERR-T01-B", skill="T01.M5")]
    links, _ = build_misconception_errors(
        [_misconception("X", "Trigger after ERR-T01-A or ERR-T01-B.")], errors)

    rows = direct_failures(links, errors)

    assert {r.micro_skill_id for r in rows} == {"T01.M2", "T01.M5"}
    assert all(r.relationship_type is RelationshipType.DIRECT_FAILURE for r in rows)


def test_two_errors_on_one_skill_produce_one_row():
    errors = [_error("ERR-T01-A", skill="T01.M2"),
              _error("ERR-T01-B", skill="T01.M2")]
    links, _ = build_misconception_errors(
        [_misconception("X", "Trigger after ERR-T01-A or ERR-T01-B.")], errors)
    assert len(direct_failures(links, errors)) == 1


def test_no_errors_means_no_skill_rows():
    assert direct_failures([], [_error("ERR-T01-A")]) == []


# ──────────────────────────────────────────────────────────────────────
# Generated: labelling wrong answers
# ──────────────────────────────────────────────────────────────────────

def _map_payload(*mappings):
    return {"mappings": list(mappings)}


def _m(qid="Q-T01-001", pattern="5n", code="ERR-T01-A"):
    return {"question_id": qid, "response_pattern": pattern, "error_code": code}


def _gen_map(payload, answers=None, questions=None, errors=None, **kw):
    answers = answers or [_answer()]
    questions = questions or [_question()]
    errors = errors or [_error("ERR-T01-A"), _error("ERR-T01-B")]
    return generate_question_error_map(
        answers, questions, errors, FakeLLMClient([payload]), "T01", **kw,
    )


def test_the_prompt_says_it_is_labelling_not_inventing():
    """CG-013 already wrote the wrong answers, checked against the accepted
    ones. A second, different set would be unreconcilable with the first."""
    text = " ".join(ERROR_MAP_SYSTEM_PROMPT.split())
    assert "You are NOT inventing wrong answers" in text


# ──────────────────────────────────────────────────────────────────────
# A letter is not a mistake
# ──────────────────────────────────────────────────────────────────────

CHOICE_TEXT = ("Which rule works? a) n + 5 b) 5n c) n - 5")


def test_option_texts_reads_what_each_option_says():
    assert option_texts(CHOICE_TEXT) == {
        "A": "n + 5", "B": "5n", "C": "n - 5",
    }


def test_option_texts_survives_options_on_their_own_lines():
    text = "Pick one.\n\na) Replace x with 5.\n\nb) Solve for x.\n\nc) Do nothing."
    assert option_texts(text)["B"] == "Solve for x."


def test_option_texts_is_empty_for_a_question_with_no_options():
    assert option_texts("Write the rule for 3+5, 9+5 using n.") == {}


def test_a_bare_letter_is_shown_with_what_the_option_says():
    """The root cause of T01's empty map on 7 September.

    69 per cent of its wrong answers were bare option letters, the highest of
    the six topics, and it returned no mappings at all. A model asked which
    error 'B' demonstrates has to go hunting for option B; shown what B says,
    it has the mistake in front of it.
    """
    prompt = build_error_map_prompt(
        [_answer(wrong="B | c")],
        [_question(text=CHOICE_TEXT)],
        [_error("ERR-T01-A")],
    )
    assert "(option B says: 5n)" in prompt
    assert "(option C says: n - 5)" in prompt


def test_a_substantive_wrong_answer_is_left_alone():
    """Only bare letters need resolving; '5n' already says what it is."""
    prompt = build_error_map_prompt(
        [_answer(wrong="5n")], [_question(text=CHOICE_TEXT)],
        [_error("ERR-T01-A")],
    )
    assert "wrong: 5n" in prompt
    assert "option" not in prompt.split("wrong: 5n")[1].split("\n")[0]


def test_the_prompt_says_a_letter_is_not_a_mistake():
    text = " ".join(ERROR_MAP_SYSTEM_PROMPT.split())
    assert "A letter is not a mistake" in text
    assert "EVERY QUESTION NEEDS AT LEAST ONE MAPPING" in text


# ──────────────────────────────────────────────────────────────────────
# Coverage: the warning that was never raised
# ──────────────────────────────────────────────────────────────────────

def test_a_partly_mapped_question_is_warned_about():
    """The gap that produced no warning at all on 7 September.

    97 questions had some wrong answers mapped and some not. The old check
    only fired when EVERY wrong answer was unmapped, so those were silent.
    """
    result = _gen_map(_map_payload(_m(pattern="5n")),
                      answers=[_answer(wrong="5n | n-5")], retry=False)
    warning = next(i for i in result[1] if "unmapped" in i.message)
    assert not warning.is_error
    assert "1 of its 2" in warning.message


def test_a_fully_mapped_question_is_not_warned_about():
    result = _gen_map(_map_payload(_m(pattern="5n"), _m(pattern="n-5")),
                      answers=[_answer(wrong="5n | n-5")], retry=False)
    assert not [i for i in result[1] if "unmapped" in i.message]


def test_unmapped_by_question_names_what_is_missing():
    rows, _, _ = _gen_map(_map_payload(_m(pattern="5n")),
                          answers=[_answer(wrong="5n | n-5")], retry=False)
    assert unmapped_by_question(rows, [_answer(wrong="5n | n-5")]) == {
        "Q-T01-001": {"n-5"},
    }


# ──────────────────────────────────────────────────────────────────────
# The retry
# ──────────────────────────────────────────────────────────────────────

def _gen_map_scripted(payloads, answers=None, questions=None, errors=None, **kw):
    answers = answers or [_answer()]
    questions = questions or [_question()]
    errors = errors or [_error("ERR-T01-A"), _error("ERR-T01-B")]
    client = FakeLLMClient(list(payloads))
    result = generate_question_error_map(
        answers, questions, errors, client, "T01", **kw,
    )
    return result, client


def test_a_question_with_nothing_is_asked_about_again():
    """One extra request is the difference between a topic whose mistakes can
    be diagnosed and T01, which had 69 questions and zero mappings."""
    (rows, issues, _), client = _gen_map_scripted(
        [_map_payload(), _map_payload(_m(pattern="5n"))],
    )
    assert len(client.calls) == 2
    assert [r.response_pattern for r in rows] == ["5n"]


def test_the_retry_asks_only_about_the_questions_that_came_back_empty():
    """Asking again about finished work wastes tokens and invites the model to
    contradict its first answer."""
    answers = [_answer(1, wrong="5n"), _answer(2, wrong="n-5")]
    questions = [_question(1), _question(2)]
    (_, _, _), client = _gen_map_scripted(
        [_map_payload(_m("Q-T01-001", "5n")),
         _map_payload(_m("Q-T01-002", "n-5"))],
        answers=answers, questions=questions,
    )
    retry_prompt = client.calls[1]["user"]
    assert "Q-T01-002" in retry_prompt
    assert "Q-T01-001" not in retry_prompt


def test_no_retry_happens_when_everything_is_mapped():
    (_, _, _), client = _gen_map_scripted([_map_payload(_m(pattern="5n"))],
                                          answers=[_answer(wrong="5n")])
    assert len(client.calls) == 1


def test_a_partly_mapped_question_does_not_trigger_a_retry():
    """It has a diagnosis for at least one mistake. Worth a warning, not
    worth a request."""
    (_, _, _), client = _gen_map_scripted(
        [_map_payload(_m(pattern="5n"))], answers=[_answer(wrong="5n | n-5")],
    )
    assert len(client.calls) == 1


def test_a_failed_retry_keeps_the_first_pass():
    """The retry exists to fill a gap. Failing to fill it must not also
    destroy what was already there."""
    answers = [_answer(1, wrong="5n"), _answer(2, wrong="n-5")]
    questions = [_question(1), _question(2)]
    (rows, issues, _), client = _gen_map_scripted(
        [_map_payload(_m("Q-T01-001", "5n")), LLMError("upstream is down")],
        answers=answers, questions=questions,
    )
    assert [r.response_pattern for r in rows] == ["5n"]
    assert any("the retry failed" in i.message for i in issues)


def test_the_retry_can_be_turned_off():
    (_, _, _), client = _gen_map_scripted([_map_payload()], retry=False)
    assert len(client.calls) == 1


def test_a_retry_that_fills_the_gap_clears_the_stale_warning():
    """A warning that a question has no diagnosis must not survive the retry
    that gave it one."""
    (rows, issues, _), _ = _gen_map_scripted(
        [_map_payload(), _map_payload(_m(pattern="5n"))],
        answers=[_answer(wrong="5n")],
    )
    assert rows
    assert not [i for i in issues
                if "no diagnosis" in i.message and i.field == "Q-T01-001"]


def test_the_prompt_states_the_job_is_labelling():
    text = " ".join(ERROR_MAP_SYSTEM_PROMPT.split())
    assert "Your only job is to say which error each one demonstrates" in text


def test_the_prompt_forbids_forcing_unrelated_mistakes_into_one_code():
    """Straight from the review."""
    text = " ".join(ERROR_MAP_SYSTEM_PROMPT.split())
    assert "Do not force unrelated mistakes into one generic error" in text


def test_the_prompt_carries_the_wrong_answers_and_the_codes():
    prompt = build_error_map_prompt(
        [_answer(1, wrong="5n | n-5")], [_question(1)],
        [_error("ERR-T01-A", name="Multiply for add")],
    )
    assert "wrong: 5n" in prompt and "wrong: n-5" in prompt
    assert "ERR-T01-A" in prompt
    assert "correct answer: n + 5" in prompt


def test_a_clean_mapping_produces_rows():
    rows, issues, _ = _gen_map(_map_payload(_m(pattern="5n"), _m(pattern="n-5")))
    assert len(rows) == 2
    assert not [i for i in issues if i.is_error]


def test_the_skill_is_filled_in_rather_than_asked_for():
    """A question has exactly one skill, so asking would only create a way to
    get it wrong."""
    rows, _, _ = _gen_map(_map_payload(_m()),
                          skill_of_question={"Q-T01-001": "T01.M4"})
    assert rows[0].micro_skill_id == "T01.M4"


def test_a_pattern_that_is_not_a_listed_wrong_answer_is_dropped():
    """A paraphrase can never match what a student actually types."""
    rows, issues, _ = _gen_map(_map_payload(_m(pattern="five n")))
    assert rows == []
    assert any("no student response will ever match it" in i.message for i in issues)


def test_an_invented_error_code_is_dropped():
    rows, issues, _ = _gen_map(
        _map_payload(_m(pattern="5n", code="ERR-T01-GHOST")))
    assert rows == []
    assert any("does not exist in this topic" in i.message for i in issues)


def test_an_unknown_question_is_dropped():
    rows, issues, _ = _gen_map(_map_payload(_m(qid="Q-T01-999")))
    assert rows == []
    assert any("not one of this topic" in i.message for i in issues)


def test_the_same_pattern_mapped_twice_is_dropped_once():
    rows, issues, _ = _gen_map(
        _map_payload(_m(pattern="5n"), _m(pattern="5n", code="ERR-T01-B")))
    assert len(rows) == 1
    assert any("twice" in i.message for i in issues)


def test_one_bad_label_does_not_lose_the_good_ones():
    """The table is additive: a wrong answer with no error attached costs one
    diagnosis, refusing the topic costs all of them."""
    rows, issues, _ = _gen_map(
        _map_payload(_m(pattern="5n"), _m(pattern="not a real answer")))
    assert len(rows) == 1
    assert all(not i.is_error for i in issues)


def test_a_question_with_nothing_mapped_warns():
    rows, issues, _ = _gen_map(_map_payload())
    assert any("gets no diagnosis" in i.message for i in issues)


def test_mapping_with_no_error_types_is_refused():
    with pytest.raises(MappingError, match="no error types"):
        generate_question_error_map(
            [_answer()], [_question()], [], FakeLLMClient([{}]), "T01")


def test_an_api_failure_propagates():
    with pytest.raises(LLMError, match="fell over"):
        generate_question_error_map(
            [_answer()], [_question()], [_error("ERR-T01-A")],
            FakeLLMClient([LLMError("the api fell over")]), "T01")


# ──────────────────────────────────────────────────────────────────────
# Generated: the relationships that are not a join
#
# DIRECT_FAILURE is derived. UNDERLYING_GAP (what the belief rests on) and
# AFFECTED_SKILL (what it spills into) are judgements, so they are asked for.
# ──────────────────────────────────────────────────────────────────────

from mapping_generator import (                     # noqa: E402
    RELATED_SKILL_SYSTEM_PROMPT,
    build_related_skills_prompt,
    generate_related_skills,
)
from models import (                                # noqa: E402
    AssessmentPriority,
    MicroSkillRow,
    MicroSkillStatus,
    MisconceptionMicroSkillRow,
)


def _skill(n):
    return MicroSkillRow(
        micro_skill_id=f"T01.M{n}", topic_id="ALG-ORI-01", skill_code=f"M{n}",
        skill_name=f"Skill {n}", description=f"The student does {n}.",
        prerequisite_micro_skill_id=None,
        assessment_priority=AssessmentPriority.HIGH,
        status=MicroSkillStatus.ACTIVE, version="1.0",
    )


def _link(mid="MIS-T01-X", skill="T01.M3", rel="UNDERLYING_GAP"):
    return {"misconception_id": mid, "micro_skill_id": skill,
            "relationship_type": rel}


def _gen_links(payload, direct=None, skills=None, misconceptions=None, **kw):
    skills = skills or [_skill(n) for n in range(1, 6)]
    misconceptions = misconceptions or [_misconception("X", "Trigger after ERR-T01-A.")]
    return generate_related_skills(
        misconceptions, skills, direct or [], FakeLLMClient([payload]), "T01", **kw,
    )


def test_the_prompt_says_direct_failures_are_already_known():
    text = " ".join(RELATED_SKILL_SYSTEM_PROMPT.split())
    assert "already known and are listed with each misconception" in text
    assert "Do not repeat them" in text


def test_the_prompt_distinguishes_cause_from_consequence():
    """If you cannot say which it is, it is probably neither."""
    text = " ".join(RELATED_SKILL_SYSTEM_PROMPT.split())
    assert "An UNDERLYING_GAP is a cause and an AFFECTED_SKILL is a consequence" in text


def test_the_prompt_warns_against_linking_everything():
    text = " ".join(RELATED_SKILL_SYSTEM_PROMPT.split())
    assert "linked to every skill in the topic says nothing useful" in text


def test_the_prompt_lists_what_is_already_broken():
    direct = [MisconceptionMicroSkillRow(
        misconception_id="MIS-T01-X", micro_skill_id="T01.M2",
        relationship_type=RelationshipType.DIRECT_FAILURE)]
    prompt = build_related_skills_prompt(
        [_misconception("X", "Trigger after ERR-T01-A.")],
        [_skill(n) for n in range(1, 4)], direct)
    assert "already recorded as directly breaking: T01.M2" in prompt


def test_a_clean_link_is_kept():
    rows, issues, _ = _gen_links({"links": [_link()]})
    assert len(rows) == 1
    assert rows[0].relationship_type is RelationshipType.UNDERLYING_GAP
    assert not [i for i in issues if i.is_error]


def test_repeating_a_direct_failure_is_dropped():
    """Recorded as broken outright; a weaker relationship on top contradicts it."""
    direct = [MisconceptionMicroSkillRow(
        misconception_id="MIS-T01-X", micro_skill_id="T01.M3",
        relationship_type=RelationshipType.DIRECT_FAILURE)]
    rows, issues, _ = _gen_links({"links": [_link(skill="T01.M3")]}, direct=direct)
    assert rows == []
    assert any("already a direct failure" in i.message for i in issues)


def test_direct_failure_cannot_be_asked_for():
    """It is derived. Accepting it here would create a second source of truth."""
    rows, issues, _ = _gen_links({"links": [_link(rel="DIRECT_FAILURE")]})
    assert rows == []
    assert any("DIRECT_FAILURE is derived, not asked for" in i.message
               for i in issues)


def test_an_invented_skill_is_dropped():
    rows, issues, _ = _gen_links({"links": [_link(skill="T01.M99")]})
    assert rows == []
    assert any("not in this topic" in i.message for i in issues)


def test_an_unknown_misconception_is_dropped():
    rows, issues, _ = _gen_links({"links": [_link(mid="MIS-T01-GHOST")]})
    assert rows == []
    assert any("does not exist" in i.message for i in issues)


def test_the_same_pair_twice_is_kept_once():
    rows, issues, _ = _gen_links(
        {"links": [_link(), _link(rel="AFFECTED_SKILL")]})
    assert len(rows) == 1
    assert any("twice" in i.message for i in issues)


def test_a_misconception_linked_to_most_of_the_topic_warns():
    """Re-teaching everything is the same as diagnosing nothing."""
    links = [_link(skill=f"T01.M{n}") for n in (1, 2, 3, 4)]
    rows, issues, _ = _gen_links({"links": links})
    assert len(rows) == 4
    assert any("does not narrow anything down" in i.message for i in issues)


def test_no_links_at_all_is_usable_but_noted():
    rows, issues, _ = _gen_links({})
    assert rows == []
    assert any("usable but thinner" in i.message for i in issues)
    assert not [i for i in issues if i.is_error]


def test_one_bad_link_does_not_lose_the_good_ones():
    rows, issues, _ = _gen_links(
        {"links": [_link(skill="T01.M2"), _link(skill="T01.M99")]})
    assert len(rows) == 1
    assert all(not i.is_error for i in issues)
