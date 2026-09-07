"""CG-019: check a written workbook end to end.

The roadmap describes this as an integration TEST on Topic 1. It is written as
a CHECKER instead, for one reason: a test proves the chain held on the day it
ran, against whatever the model happened to return. A checker can be pointed
at any workbook, including the ones actually sent for review, and answers the
question that matters -- does this file hold together?

Everything here is driven by TABLE_SCHEMAS rather than written out by hand.
The schema already declares 34 foreign keys and 20 unique columns, so a table
added later is checked without anyone remembering to add it. A hand-written
list of joins is a list that goes stale the first time someone is in a hurry.

What it checks
--------------

    references     every foreign key resolves to a row that exists
    uniqueness     no id is issued twice
    orphans        rows nothing points at, where that means something is lost
    coverage       every micro-skill has its seven questions
    the review     one mapping per question, primary, weight 1.0

What it does not check
-----------------------

Whether the mathematics is right. That needs a person or an independent model,
and it is CG-021's job. A workbook can pass every check here and still teach a
student something false, which is exactly why nothing this pipeline produces is
marked APPROVED.
"""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from coverage_plan import coverage_report, plan_for_skill
from models import Phase
from table_schemas import TABLE_SCHEMAS, export_name

#: Tables that are legitimately empty because nothing generates them.
NOT_GENERATED = {
    "Orientation_Videos", "Orientation_Video_Scenes", "Orientation_Support_Cards",
}


@dataclass(frozen=True)
class Problem:
    """One thing wrong with a workbook."""

    kind: str
    table: str
    detail: str

    def __str__(self) -> str:
        return f"[{self.kind}] {self.table}: {self.detail}"


def read_tables(path: str | Path) -> dict[str, list[dict]]:
    """Every sheet as rows, keyed by the INTERNAL table name."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        by_export = {export_name(name): name for name in TABLE_SCHEMAS}
        tables: dict[str, list[dict]] = {}
        for sheet in workbook.sheetnames:
            internal = by_export.get(sheet)
            if internal is None:
                continue
            worksheet = workbook[sheet]
            rows = worksheet.iter_rows(values_only=True)
            header = [h for h in (next(rows, ()) or ()) if h is not None]
            tables[internal] = [
                dict(zip(header, values)) for values in rows
                if any(v is not None for v in values)
            ]
        return tables
    finally:
        workbook.close()


def check_references(tables: dict[str, list[dict]]) -> list[Problem]:
    """Every foreign key points at a row that exists.

    Driven by the schema's own declarations, so a table added later is checked
    without anyone remembering to add it here.
    """
    problems: list[Problem] = []

    for table, schema in TABLE_SCHEMAS.items():
        rows = tables.get(table, [])
        for column, (target_table, target_column) in (
            schema.get("foreign_keys") or {}
        ).items():
            known = {
                row.get(target_column) for row in tables.get(target_table, [])
                if row.get(target_column) is not None
            }
            dangling = Counter(
                row[column] for row in rows
                if row.get(column) is not None and row[column] not in known
            )
            for value, count in dangling.items():
                problems.append(Problem(
                    "reference", table,
                    f"{column}={value!r} does not exist in "
                    f"{target_table}.{target_column}"
                    + (f" ({count} rows)" if count > 1 else ""),
                ))
    return problems


def check_uniqueness(tables: dict[str, list[dict]]) -> list[Problem]:
    """No id issued twice.

    A duplicate id is worse than a missing one: a reference to it resolves,
    silently, to whichever row is found first.
    """
    problems: list[Problem] = []

    for table, schema in TABLE_SCHEMAS.items():
        rows = tables.get(table, [])
        for column, details in schema["column_details"].items():
            if not details.get("unique"):
                continue
            counts = Counter(
                row[column] for row in rows if row.get(column) is not None)
            for value, count in counts.items():
                if count > 1:
                    problems.append(Problem(
                        "duplicate", table,
                        f"{column}={value!r} appears {count} times",
                    ))
    return problems


def check_orphans(tables: dict[str, list[dict]]) -> list[Problem]:
    """Rows nothing points at, where that means something is unreachable.

    Not every unreferenced row is a fault -- a topic is referenced but never
    references anything -- so only the cases where being unreachable makes the
    row useless are listed.
    """
    problems: list[Problem] = []

    # Which table has to point at it, not merely SOME table. A question with
    # a micro-skill mapping but no usage row is still never shown to anyone,
    # so "is anything referencing this?" would call it reachable and be wrong.
    # Reachability is a different question per table, so each says its own.
    unreachable = [
        # (table, id column, table that must reference it, its column, why)
        ("Questions", "question_id", "Question_Usage", "question_id",
         "a question in no usage row is never shown to anyone"),
        ("Hints", "hint_id", "Misconception_Hints", "hint_id",
         "a hint linked to no misconception can never be offered"),
        ("Visual_Cues", "visual_cue_id", "Misconception_VisualCues",
         "visual_cue_id",
         "a cue linked to no misconception can never be shown"),
        ("Scaffolds", "scaffold_id", "Question_Scaffolds", "scaffold_id",
         "a scaffold attached to no question can never start"),
    ]

    for table, column, by_table, by_column, why in unreachable:
        pointed_at = {
            row[by_column] for row in tables.get(by_table, [])
            if row.get(by_column) is not None
        }
        for row in tables.get(table, []):
            value = row.get(column)
            if value is not None and value not in pointed_at:
                problems.append(Problem("orphan", table, f"{value}: {why}"))
    return problems


def check_review_rules(tables: dict[str, list[dict]]) -> list[Problem]:
    """The mapping rules the content review made explicit.

    Exactly one micro-skill per question, primary, at weight 1.0. These are
    checked on the written file rather than trusted from the generator,
    because the file is what the platform imports.
    """
    problems: list[Problem] = []
    mappings = tables.get("Question_MicroSkills", [])
    questions = tables.get("Questions", [])

    per_question = Counter(r["question_id"] for r in mappings
                           if r.get("question_id"))
    for question_id, count in per_question.items():
        if count != 1:
            problems.append(Problem(
                "review", "Question_MicroSkills",
                f"{question_id} has {count} micro-skill mappings; the review "
                f"requires exactly one",
            ))

    unmapped = {r["question_id"] for r in questions if r.get("question_id")} \
        - set(per_question)
    for question_id in sorted(unmapped):
        problems.append(Problem(
            "review", "Question_MicroSkills",
            f"{question_id} has no micro-skill mapping, so answering it "
            f"tells the student model nothing",
        ))

    for row in mappings:
        if row.get("is_primary") is not True:
            problems.append(Problem(
                "review", "Question_MicroSkills",
                f"{row.get('question_id')} -> {row.get('micro_skill_id')} is "
                f"not primary; the review requires is_primary TRUE",
            ))
        if row.get("weight") != 1:
            problems.append(Problem(
                "review", "Question_MicroSkills",
                f"{row.get('question_id')} -> {row.get('micro_skill_id')} has "
                f"weight {row.get('weight')}; the review requires 1.0",
            ))

    # Nothing generated may claim to have been reviewed.
    approved = [r for r in questions if r.get("status") == "APPROVED"]
    if approved:
        problems.append(Problem(
            "review", "Questions",
            f"{len(approved)} question(s) marked APPROVED; nothing may be "
            f"approved until the CG-020 and CG-021 checks have run",
        ))
    return problems


def check_coverage(tables: dict[str, list[dict]]) -> list[Problem]:
    """Every micro-skill has the seven questions the plan requires."""
    usage = {r["question_id"]: r for r in tables.get("Question_Usage", [])
             if r.get("question_id")}
    difficulty = {r["question_id"]: r.get("difficulty")
                  for r in tables.get("Questions", []) if r.get("question_id")}

    by_skill: dict[str, list[tuple[Phase, int]]] = {
        row["micro_skill_id"]: []
        for row in tables.get("Micro_Skills", []) if row.get("micro_skill_id")
    }
    for row in tables.get("Question_MicroSkills", []):
        question_id = row.get("question_id")
        skill_id = row.get("micro_skill_id")
        if skill_id not in by_skill or question_id not in usage:
            continue
        try:
            phase = Phase(usage[question_id]["phase"])
        except ValueError:
            continue
        by_skill[skill_id].append((phase, difficulty.get(question_id)))

    return [
        Problem("coverage", "Questions",
                f"{skill_id} is missing {', '.join(str(s) for s in gaps)}")
        for skill_id, gaps in coverage_report(by_skill).items()
    ]


def check_workbook(path: str | Path) -> list[Problem]:
    """Everything, in one pass. An empty list means the file holds together."""
    tables = read_tables(path)
    return (
        check_references(tables)
        + check_uniqueness(tables)
        + check_orphans(tables)
        + check_review_rules(tables)
        + check_coverage(tables)
    )


def summarise(problems: list[Problem], path: Optional[str | Path] = None) -> str:
    """A short report of what a check found."""
    header = f"{Path(path).name}: " if path else ""
    if not problems:
        return (f"{header}every reference resolves, every id is unique, and "
                f"every micro-skill has its questions.")

    counts = Counter(p.kind for p in problems)
    lines = [
        f"{header}{len(problems)} problem(s): "
        + ", ".join(f"{n} {kind}" for kind, n in counts.most_common()),
        "",
    ]
    lines += [f"  {p}" for p in problems[:40]]
    if len(problems) > 40:
        lines.append(f"  ... and {len(problems) - 40} more")
    return "\n".join(lines)


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("usage: python integrity.py <workbook.xlsx>")
        raise SystemExit(2)

    found = check_workbook(sys.argv[1])
    print(summarise(found, sys.argv[1]))
    raise SystemExit(1 if found else 0)
